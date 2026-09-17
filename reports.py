import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from datetime import datetime

from database import get_connection
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment


class ReportsWindow:

    def __init__(self, parent, current_user):

        self.current_user = current_user

        self.window = tk.Toplevel(parent)
        self.window.title("Reports")
        self.window.geometry("900x550")
        self.window.resizable(True, True)

        self.build_interface()
        self.load_report()

    def build_interface(self):

        tk.Label(
            self.window,
            text="SYSTEM REPORT",
            font=("Segoe UI", 18, "bold")
        ).pack(pady=20)

        table_frame = tk.Frame(self.window)
        table_frame.pack(
            fill="both",
            expand=True,
            padx=30,
            pady=10
        )

        self.table = ttk.Treeview(
            table_frame,
            columns=("Item", "Count"),
            show="headings"
        )

        self.table.heading(
            "Item",
            text="Item"
        )

        self.table.heading(
            "Count",
            text="Count"
        )

        self.table.column(
            "Item",
            width=350,
            anchor="w"
        )

        self.table.column(
            "Count",
            width=200,
            anchor="center"
        )

        scrollbar = ttk.Scrollbar(
            table_frame,
            orient="vertical",
            command=self.table.yview
        )

        self.table.configure(
            yscrollcommand=scrollbar.set
        )

        self.table.pack(
            side="left",
            fill="both",
            expand=True
        )

        scrollbar.pack(
            side="right",
            fill="y"
        )

        # Buttons
        button_frame = tk.Frame(self.window)
        button_frame.pack(pady=15)

        ttk.Button(
            button_frame,
            text="Refresh",
            command=self.load_report
        ).pack(
            side="left",
            padx=5
        )

        ttk.Button(
            button_frame,
            text="Export to Excel",
            command=self.export_to_excel
        ).pack(
            side="left",
            padx=5
        )

        ttk.Button(
            button_frame,
            text="Close",
            command=self.window.destroy
        ).pack(
            side="left",
            padx=5
        )

    def get_report_data(self):

        connection = get_connection()

        if not connection:
            raise Exception(
                "Could not connect to the database."
            )

        cursor = connection.cursor()

        queries = [
            (
                "Total Users",
                "SELECT COUNT(*) FROM users"
            ),
            (
                "Total Devices",
                "SELECT COUNT(*) FROM devices"
            ),
            (
                "Pending Requests",
                """
                SELECT COUNT(*)
                FROM access_requests
                WHERE status = 'PENDING'
                """
            ),
            (
                "Approved Requests",
                """
                SELECT COUNT(*)
                FROM access_requests
                WHERE status = 'APPROVED'
                """
            ),
            (
                "Rejected Requests",
                """
                SELECT COUNT(*)
                FROM access_requests
                WHERE status = 'REJECTED'
                """
            ),
            (
                "Remote Sessions",
                """
                SELECT COUNT(*)
                FROM remote_sessions
                """
            )
        ]

        report_data = []

        try:

            for name, query in queries:

                cursor.execute(query)

                result = cursor.fetchone()

                count = result[0] if result else 0

                report_data.append(
                    (name, count)
                )

            return report_data

        finally:

            cursor.close()
            connection.close()

    def load_report(self):

        # Clear old report
        for item in self.table.get_children():
            self.table.delete(item)

        try:

            report_data = self.get_report_data()

            for name, count in report_data:

                self.table.insert(
                    "",
                    tk.END,
                    values=(name, count)
                )

        except Exception as error:

            messagebox.showerror(
                "Report Error",
                str(error),
                parent=self.window
            )

    def export_to_excel(self):

        try:

            # Get current data from MySQL
            report_data = self.get_report_data()

            # Ask user where to save the Excel file
            file_path = filedialog.asksaveasfilename(
                parent=self.window,
                title="Save Excel Report",
                defaultextension=".xlsx",
                filetypes=[
                    (
                        "Excel Workbook",
                        "*.xlsx"
                    )
                ],
                initialfile=(
                    "Remote_Access_System_Report_"
                    + datetime.now().strftime(
                        "%Y%m%d_%H%M%S"
                    )
                    + ".xlsx"
                )
            )

            # User cancelled
            if not file_path:
                return

            # Create Excel workbook
            workbook = Workbook()

            worksheet = workbook.active
            worksheet.title = "System Report"

            # Report title
            worksheet["A1"] = (
                "REMOTE ACCESS MANAGEMENT SYSTEM"
            )

            worksheet["A1"].font = Font(
                bold=True,
                size=16
            )

            # Report subtitle
            worksheet["A2"] = "SYSTEM REPORT"

            worksheet["A2"].font = Font(
                bold=True,
                size=14
            )

            # Date and time
            worksheet["A3"] = "Generated"
            worksheet["B3"] = datetime.now().strftime(
                "%Y-%m-%d %H:%M:%S"
            )

            # Table headings
            worksheet["A5"] = "Item"
            worksheet["B5"] = "Count"

            worksheet["A5"].font = Font(
                bold=True
            )

            worksheet["B5"].font = Font(
                bold=True
            )

            worksheet["A5"].alignment = Alignment(
                horizontal="center"
            )

            worksheet["B5"].alignment = Alignment(
                horizontal="center"
            )

            # Add report data
            row_number = 6

            for name, count in report_data:

                worksheet.cell(
                    row=row_number,
                    column=1,
                    value=name
                )

                worksheet.cell(
                    row=row_number,
                    column=2,
                    value=count
                )

                row_number += 1

            # Adjust column widths
            worksheet.column_dimensions[
                "A"
            ].width = 35

            worksheet.column_dimensions[
                "B"
            ].width = 20

            # Save Excel file
            workbook.save(file_path)

            messagebox.showinfo(
                "Export Successful",
                "Excel report generated successfully.\n\n"
                f"Saved to:\n{file_path}",
                parent=self.window
            )

        except Exception as error:

            messagebox.showerror(
                "Export Error",
                str(error),
                parent=self.window
            )